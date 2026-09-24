import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import tkinter as tk
from tkinter import filedialog, messagebox
from pathlib import Path
from datetime import datetime
import re


# --------- Helpers ----------
def clean_column_names(df):
    df.columns = [
        str(col)
        .replace('\xa0', ' ')
        .replace('\u200b', '')
        .strip()
        for col in df.columns
    ]
    return df


def normalize_value(val):
    if pd.isna(val):
        return ""
    return str(val).strip().lower()


def to_number(series):
    """Convert '$242,000.00' or '242000' strings to plain numeric strings."""
    return (
        series.astype(str)
        .str.replace("$", "", regex=False)
        .str.replace(",", "", regex=False)
        .str.strip()
    )


# --------- Core comparison ----------
def compare_salary_files(pre_file, post_file):
    # Pre: normal header at row 1
    pre = pd.read_excel(pre_file, dtype=str)
    # Post: header on row 4 (index 3)
    post = pd.read_excel(post_file, dtype=str, header=3)

    pre = clean_column_names(pre)
    post = clean_column_names(post)

    # helper finders
    def get_emp_id_col(cols):
        for name in ["Employee ID", "Employee Id", "Emp ID", "Emp Id"]:
            if name in cols:
                return name
        return None

    def get_ssn_col(cols):
        for name in ["SSN", "SSN#", "Social Security Number"]:
            if name in cols:
                return name
        return None

    pre_emp_col = get_emp_id_col(pre.columns)
    post_emp_col = get_emp_id_col(post.columns)
    pre_ssn_col = get_ssn_col(pre.columns)
    post_ssn_col = get_ssn_col(post.columns)

    if not post_emp_col:
        raise ValueError(
            "Post file must contain Employee ID / Employee Id column.\n\n"
            f"Post columns: {list(post.columns)}"
        )

    # standardize keys
    post = post.rename(columns={post_emp_col: "Employee ID"})
    post["Employee ID"] = post["Employee ID"].astype(str)

    if post_ssn_col:
        post = post.rename(columns={post_ssn_col: "SSN"})
        post["SSN"] = post["SSN"].astype(str)

    if pre_emp_col:
        pre = pre.rename(columns={pre_emp_col: "Employee ID"})
        pre["Employee ID"] = pre["Employee ID"].astype(str)
    elif pre_ssn_col:
        pre = pre.rename(columns={pre_ssn_col: "SSN"})
        pre["SSN"] = pre["SSN"].astype(str)
    else:
        raise ValueError(
            "Pre file must contain either Employee ID / Employee Id or SSN/SSN#.\n\n"
            f"Pre columns: {list(pre.columns)}"
        )

    # salary columns
    if "Annual Salary" not in pre.columns:
        raise ValueError(
            f"Pre file must contain 'Annual Salary' column. Found: {list(pre.columns)}"
        )
    if "Salary" not in post.columns:
        raise ValueError(
            f"Post file must contain 'Salary' column. Found: {list(post.columns)}"
        )

    PRE_SAL_COL = "Annual Salary"
    POST_SAL_COL = "Salary"
    WORK_STATUS_COL = "Status"
    WORK_SALPERIOD_COL = "Salary Period"

    pre[PRE_SAL_COL] = to_number(pre[PRE_SAL_COL])
    post[POST_SAL_COL] = to_number(post[POST_SAL_COL])

    post_lookup = post.set_index("Employee ID")
    post_lookup_ssn = post.set_index("SSN") if "SSN" in post.columns else None

    comparison_salary = []
    salary_from_benj = []
    status_from_benj = []
    salperiod_from_benj = []
    diff_values = []
    results = []
    comments = []

    for _, row in pre.iterrows():
        pre_sal = row[PRE_SAL_COL]
        post_row = None

        if "Employee ID" in pre.columns:
            emp_id = str(row["Employee ID"])
            if emp_id in post_lookup.index:
                post_row = post_lookup.loc[emp_id]
                if isinstance(post_row, pd.DataFrame):
                    post_row = post_row.iloc[0]

        if post_row is None and "SSN" in pre.columns and post_lookup_ssn is not None:
            ssn_val = str(row["SSN"])
            if ssn_val in post_lookup_ssn.index:
                post_row = post_lookup_ssn.loc[ssn_val]
                if isinstance(post_row, pd.DataFrame):
                    post_row = post_row.iloc[0]

        if post_row is not None:
            post_sal = post_row[POST_SAL_COL]
            if isinstance(post_sal, pd.Series):
                post_sal = post_sal.iloc[0]

            work_status = post_row.get(WORK_STATUS_COL, "")
            work_salperiod = post_row.get(WORK_SALPERIOD_COL, "")
            if isinstance(work_status, pd.Series):
                work_status = work_status.iloc[0]
            if isinstance(work_salperiod, pd.Series):
                work_salperiod = work_salperiod.iloc[0]

            comparison_salary.append(pre_sal)            # from Salary File
            salary_from_benj.append(post_sal)            # from BenJ Work File
            status_from_benj.append(work_status)
            salperiod_from_benj.append(work_salperiod)

            try:
                pre_num = float(pre_sal)
                post_num = float(post_sal)
                same = round(pre_num, 2) == round(post_num, 2)
                diff_values.append(round(post_num - pre_num, 2))
            except ValueError:
                same = normalize_value(pre_sal) == normalize_value(post_sal)
                diff_values.append("")

            results.append(same)
            comments.append("" if same else "Salary changed")
        else:
            comparison_salary.append("")
            salary_from_benj.append("")
            status_from_benj.append("")
            salperiod_from_benj.append("")
            diff_values.append("")
            results.append(False)
            comments.append("New Employee (not in post file)")

    result_df = pre.copy()
    result_df["Salary From Salary File"] = comparison_salary
    result_df["Salary From WorkFile"] = salary_from_benj
    result_df["Status From WorkFile"] = status_from_benj
    result_df["Salary Period From Benj"] = salperiod_from_benj
    result_df["Diffrence"] = diff_values
    result_df["Results"] = results
    result_df["Comment"] = comments

    downloads = Path.home() / "Downloads"
    downloads.mkdir(exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = downloads / f"Salary_Comparison_{timestamp}.xlsx"
    result_df.to_excel(output_file, index=False)

    wb = load_workbook(output_file)
    ws = wb.active
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    date_col_indices = set()
    for col_idx in range(1, ws.max_column + 1):
        header_val = str(ws.cell(row=1, column=col_idx).value or "").strip().lower()
        if "compare" not in header_val and ("date" in header_val or "dob" in header_val or "birth" in header_val or "hire" in header_val or "term" in header_val or "effective" in header_val):
            date_col_indices.add(col_idx)

    for col_idx in range(1, ws.max_column + 1):
        is_date_col = col_idx in date_col_indices
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is None:
                continue
            if is_date_col or (isinstance(val, str) and re.search(r"^\d{4}-\d{2}-\d{2}", val.strip())):
                try:
                    ts = pd.to_datetime(str(val).strip(), errors='coerce')
                    if pd.notna(ts):
                        cell.value = ts.date()
                        cell.number_format = "MM/DD/YYYY"
                except Exception:
                    pass

    for cell in ws[1]:
        if cell.value in [
            "Salary From Salary File",
            "Salary From WorkFile",
            "Status From WorkFile",
            "Salary Period From Benj",
            "Diffrence",
            "Results",
            "Comment",
        ]:
            cell.fill = yellow
    wb.save(output_file)

    return str(output_file), result_df


# --------- Fill Template (Employee ID or SSN/Country ID) ----------
def fill_template_file(comparison_file, template_file):
    """
    From comparison file:
      - Take rows where Status From WorkFile = Active AND Results = False
      - For each, use Employee ID if present, else SSN (Country ID in EmployeeInfo)
      - Write Employee Id / SSN / Salary / Salary Period into template
    """
    comparison_df = pd.read_excel(comparison_file, dtype=str)
    comparison_df = clean_column_names(comparison_df)

    active_false = comparison_df[
        (comparison_df["Status From WorkFile"].astype(str).str.strip().str.lower() == "active") &
        (comparison_df["Results"].astype(str).str.strip().str.lower() == "false")
    ].copy()

    if active_false.empty:
        messagebox.showwarning("Info", "No records found with Status=Active AND Results=FALSE")
        return

    template_df = pd.read_excel(template_file, dtype=str)
    template_df = clean_column_names(template_df)

    # NEW: In EmployeeInfo, 'Country ID' is the SSN used for matching
    if "Country ID" in template_df.columns and "SSN" not in template_df.columns:
        template_df = template_df.rename(columns={"Country ID": "SSN"})

    # detect template columns
    emp_id_col = None
    for name in ["Employee Id", "Employee ID", "Emp ID", "Emp Id"]:
        if name in template_df.columns:
            emp_id_col = name
            break

    ssn_col = None
    for name in ["SSN", "Ssn", "Social Security Number"]:
        if name in template_df.columns:
            ssn_col = name
            break

    sal_col = None
    for name in ["Salary"]:
        if name in template_df.columns:
            sal_col = name
            break

    sal_period_col = None
    for name in ["Salary Period"]:
        if name in template_df.columns:
            sal_period_col = name
            break

    if not sal_col or not sal_period_col:
        raise ValueError(
            f"Template must have at least 'Salary' and 'Salary Period' columns.\n"
            f"Found: {list(template_df.columns)}"
        )

    rename_map = {}
    if emp_id_col:
        rename_map[emp_id_col] = "Employee Id"
    if ssn_col:
        rename_map[ssn_col] = "SSN"
    rename_map[sal_col] = "Salary"
    rename_map[sal_period_col] = "Salary Period"
    template_df = template_df.rename(columns=rename_map)

    if "Employee Id" not in template_df.columns:
        template_df["Employee Id"] = ""
    if "SSN" not in template_df.columns and "SSN" in comparison_df.columns:
        template_df["SSN"] = ""

    filled_count = 0

    for _, comp_row in active_false.iterrows():
        emp_id_val = str(comp_row.get("Employee ID", "")).strip()
        ssn_val = str(comp_row.get("SSN", "")).strip()

        if not emp_id_val and not ssn_val:
            continue

        # UPDATED: salary from Salary File, not Work File
        salary = str(comp_row.get("Salary From Salary File", "")).strip()
        sal_period = str(comp_row.get("Salary Period From Benj", "")).strip()

        empty_mask = (
            template_df["Employee Id"].astype(str).str.strip().eq("") &
            template_df["SSN"].astype(str).str.strip().eq("")
        )

        if empty_mask.any():
            idx = empty_mask.idxmax()
        else:
            new_row = {col: "" for col in template_df.columns}
            template_df = pd.concat([template_df, pd.DataFrame([new_row])], ignore_index=True)
            idx = template_df.index[-1]

        if emp_id_val:
            template_df.at[idx, "Employee Id"] = emp_id_val
        if ssn_val:
            template_df.at[idx, "SSN"] = ssn_val
        template_df.at[idx, "Salary"] = salary
        template_df.at[idx, "Salary Period"] = sal_period
        filled_count += 1

    output_file = Path(template_file).parent / (
        f"{Path(template_file).stem}_Filled_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    template_df.to_excel(output_file, index=False)

    wb = load_workbook(output_file)
    ws = wb.active
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    for cell in ws[1]:
        if cell.value in ["Employee Id", "SSN", "Salary", "Salary Period"]:
            col_idx = cell.column
            cell.fill = yellow
            for row_idx in range(2, ws.max_row + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                if c.value and str(c.value).strip():
                    c.fill = yellow

    for col_idx in range(1, ws.max_column + 1):
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None and isinstance(val, str) and re.search(r"^\d{4}-\d{2}-\d{2}", val.strip()):
                try:
                    ts = pd.to_datetime(str(val).strip(), errors='coerce')
                    if pd.notna(ts):
                        cell.value = ts.date()
                        cell.number_format = "MM/DD/YYYY"
                except Exception:
                    pass

    wb.save(output_file)

    messagebox.showinfo(
        "Success ✅",
        f"Template filled successfully!\n\n"
        f"Records filled: {filled_count}\n"
        f"Output saved at:\n{output_file}"
    )

    return str(output_file)


# --------- GUI ----------
def run_gui():
    root = tk.Tk()
    root.title("Salary Comparison & Upload Template Tool")
    root.geometry("750x600")
    root.config(bg="#f6f6f6")

    tk.Label(root, text="STEP 1: SALARY COMPARISON", bg="#f6f6f6",
             font=("Arial", 12, "bold", "underline")).pack(pady=10)

    tk.Label(root, text="Salary File:", bg="#f6f6f6",
             font=("Arial", 10)).pack(pady=3)
    pre_entry = tk.Entry(root, width=90)
    pre_entry.pack()
    tk.Button(
        root, text="Browse", bg="#4CAF50", fg="white",
        command=lambda: pre_entry.delete(0, tk.END) or pre_entry.insert(
            0, filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        )
    ).pack(pady=4)

    tk.Label(root, text="Work File from BenJ:", bg="#f6f6f6",
             font=("Arial", 10)).pack(pady=3)
    post_entry = tk.Entry(root, width=90)
    post_entry.pack()
    tk.Button(
        root, text="Browse", bg="#4CAF50", fg="white",
        command=lambda: post_entry.delete(0, tk.END) or post_entry.insert(
            0, filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        )
    ).pack(pady=4)

    tk.Label(root, text="STEP 2: UPLOAD TEMPLATE", bg="#f6f6f6",
             font=("Arial", 12, "bold", "underline")).pack(pady=15)

    tk.Label(root, text="Upload Template File:", bg="#f6f6f6",
             font=("Arial", 10)).pack(pady=3)
    template_entry = tk.Entry(root, width=90)
    template_entry.pack()
    tk.Button(
        root, text="Browse", bg="#4CAF50", fg="white",
        command=lambda: template_entry.delete(0, tk.END) or template_entry.insert(
            0, filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")])
        )
    ).pack(pady=4)

    def generate_and_fill():
        pre_file = pre_entry.get().strip()
        post_file = post_entry.get().strip()
        template_file = template_entry.get().strip()

        if not pre_file or not post_file or not template_file:
            messagebox.showwarning("Warning", "Please select all three files")
            return

        try:
            comparison_path, _ = compare_salary_files(pre_file, post_file)
            messagebox.showinfo(
                "Step 1 Complete ✅",
                f"Salary comparison completed!\n\nOutput saved at:\n{comparison_path}"
            )
            fill_template_file(comparison_path, template_file)
        except Exception as e:
            messagebox.showerror("Error", str(e))

    tk.Button(
        root, text="Generate & Fill Template",
        bg="#2196F3", fg="white",
        font=("Arial", 12, "bold"),
        command=generate_and_fill
    ).pack(pady=25)

    tk.Label(root, text="Process: Salary Comparison → Filter (Active + FALSE) → Fill Template",
             bg="#f6f6f6", font=("Arial", 9, "italic")).pack(pady=10)

    root.mainloop()


if __name__ == "__main__":
    run_gui()
