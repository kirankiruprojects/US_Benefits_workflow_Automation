import re
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import tkinter as tk
from tkinter import filedialog, messagebox

# ---------------- Helpers ----------------
def clean_column_names(df):
    df.columns = [str(col).replace('\xa0', ' ').replace('\u200b', '').strip() for col in df.columns]
    return df

def normalize_ssn(ssn):
    if pd.isna(ssn):
        return ""
    return "".join(filter(str.isalnum, str(ssn))).lower()

def normalize_text(val):
    if pd.isna(val):
        return ""
    return str(val).replace('\xa0', ' ').replace('\u200b', '').strip().lower()

def get_status_column(df):
    possible = ["Enrollment Status", "Employment Status", "Status"]
    for c in df.columns:
        if str(c).strip().lower() in [p.lower() for p in possible]:
            return c
    raise ValueError("Status column not found")

def is_sp_or_ch(name):
    low = name.lower()
    return bool(re.search(r'\b(sp|ch)\b', low))

def choose_er_pays_column(cols):
    if not cols:
        return None
    filtered = [c for c in cols if not is_sp_or_ch(c)]
    if filtered:
        cols = filtered
    for c in cols:
        if re.search(r'\bee\b', c, re.I) or re.search(r'employee', c, re.I):
            return c
    return cols[0]

def choose_volume_column(cols):
    if not cols:
        return None
    filtered = [c for c in cols if not is_sp_or_ch(c)]
    if filtered:
        return filtered[0]
    return cols[0]

# ---------------- Main Logic ----------------
def generate_output(pre_file, post_file, output_file):
    pre = pd.read_excel(pre_file, header=4, dtype=str)
    post = pd.read_excel(post_file, header=4, dtype=str)

    pre = clean_column_names(pre)
    post = clean_column_names(post)

    status_pre = get_status_column(pre)
    status_post = get_status_column(post)

    pre_active = pre[pre[status_pre].str.strip().str.lower() == "active"].copy()
    post_active = post[post[status_post].str.strip().str.lower() == "active"].copy()

    pre_active["ssn_norm"] = pre_active["Employee SSN"].map(normalize_ssn)
    post_active["ssn_norm"] = post_active["Employee SSN"].map(normalize_ssn)

    plan_name_cols = [
        col for col in post_active.columns
        if str(col).strip().lower().endswith("plan name")
    ]

    if not plan_name_cols:
        raise ValueError("No plan columns ending with 'Plan Name' found in Post file")

    excluded_prefixes = ("medical", "dental", "vision")
    plan_name_cols = [p for p in plan_name_cols if not p.strip().lower().startswith(excluded_prefixes)]

    basic_suffixes = ["Plan Name", "Effective Date", "Coverage"]

    for plan_col in plan_name_cols:
        plan_prefix = plan_col.replace("Plan Name", "").strip()

        plan_columns = []

        for suf in basic_suffixes:
            col_name = f"{plan_prefix} {suf}" if suf != "Plan Name" else plan_col
            if col_name in post_active.columns:
                plan_columns.append(col_name)

        er_candidates = [
            c for c in post_active.columns
            if c.lower().startswith(plan_prefix.lower()) and c.lower().strip().endswith("er pays")
        ]
        er_col = choose_er_pays_column(er_candidates)
        if er_col:
            plan_columns.append(er_col)

        vol_candidates = [
            c for c in post_active.columns
            if c.lower().startswith(plan_prefix.lower()) and c.lower().strip().endswith("volume")
        ]
        vol_col = choose_volume_column(vol_candidates)
        if vol_col:
            plan_columns.append(vol_col)

        if not plan_columns:
            continue

        for col in plan_columns:
            original_post_series = post_active[col].copy()
            if col in pre_active.columns:
                pre_lookup = dict(zip(pre_active["ssn_norm"], pre_active[col]))
            else:
                pre_lookup = {}

            pre_values = []
            compare_results = []

            for idx in post_active.index:
                ssn_norm = post_active.at[idx, "ssn_norm"]
                post_val = original_post_series.at[idx]

                if ssn_norm and ssn_norm in pre_lookup:
                    pre_val = pre_lookup[ssn_norm]
                    pre_values.append(pre_val if pd.notna(pre_val) else "#N/A")
                    compare_results.append(normalize_text(post_val) == normalize_text(pre_val))
                else:
                    pre_values.append("#N/A")
                    compare_results.append("#N/A")

            post_active[col] = pre_values

            insert_idx = post_active.columns.get_loc(col) + 1
            post_col_name = f"{col} (Post)"
            if post_col_name in post_active.columns:
                i = 1
                while f"{post_col_name} ({i})" in post_active.columns:
                    i += 1
                post_col_name = f"{post_col_name} ({i})"

            post_active.insert(insert_idx, post_col_name, original_post_series.values)

            compare_col_name = f"{col} Compare"
            if compare_col_name in post_active.columns:
                i = 1
                while f"{compare_col_name} ({i})" in post_active.columns:
                    i += 1
                compare_col_name = f"{compare_col_name} ({i})"

            post_active.insert(insert_idx + 1, compare_col_name, compare_results)

    import datetime as dt
    for c in post_active.columns:
        c_low = str(c).strip().lower()
        if "compare" not in c_low and ("date" in c_low or "dob" in c_low or "birth" in c_low or "hire" in c_low or "term" in c_low or "effective" in c_low):
            try:
                post_active[c] = pd.to_datetime(post_active[c], errors='coerce').dt.strftime('%m/%d/%Y')
            except Exception:
                pass

    post_active.to_excel(output_file, index=False)

    wb = load_workbook(output_file)
    ws = wb.active
    yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    date_col_indices = set()
    for col_idx in range(1, ws.max_column + 1):
        header_val = str(ws.cell(row=1, column=col_idx).value or "").strip().lower()
        if "compare" not in header_val and ("date" in header_val or "dob" in header_val or "birth" in header_val or "hire" in header_val or "term" in header_val or "effective" in header_val):
            date_col_indices.add(col_idx)

    for col_idx in range(1, ws.max_column + 1):
        is_date_col = col_idx in date_col_indices
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is None:
                continue
            if is_date_col or (isinstance(val, str) and re.search(r"^\d{4}-\d{2}-\d{2}", val.strip())):
                try:
                    ts = pd.to_datetime(str(val).strip(), errors='coerce')
                    if pd.notna(ts):
                        cell.value = ts.date()
                        cell.number_format = "MM/DD/YYYY"
                except Exception:
                    pass

    for cell in ws[1]:
        if cell.value and ("(Post)" in str(cell.value) or " Compare" in str(cell.value)):
            cell.fill = yellow

    wb.save(output_file)

# ---------------- GUI ----------------
def run_gui():
    root = tk.Tk()
    root.title("Pre → Post Excel Comparison Tool")
    root.geometry("640x360")
    root.config(bg="#f6f6f6")

    tk.Label(root, text="Pre Excel File:", bg="#f6f6f6", font=("Arial", 11)).pack(pady=6)
    pre_entry = tk.Entry(root, width=80)
    pre_entry.pack()
    tk.Button(root, text="Browse", bg="#4CAF50", fg="white",
              command=lambda: pre_entry.delete(0, tk.END) or pre_entry.insert(
                  0, filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")]))).pack(pady=6)

    tk.Label(root, text="Post Excel File:", bg="#f6f6f6", font=("Arial", 11)).pack(pady=6)
    post_entry = tk.Entry(root, width=80)
    post_entry.pack()
    tk.Button(root, text="Browse", bg="#4CAF50", fg="white",
              command=lambda: post_entry.delete(0, tk.END) or post_entry.insert(
                  0, filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")]))).pack(pady=6)

    def generate():
        pre_file = pre_entry.get().strip()
        post_file = post_entry.get().strip()
        if not pre_file or not post_file:
            messagebox.showwarning("Warning", "Please select both Pre and Post files.")
            return

        # Save automatically to Downloads folder
        downloads_folder = os.path.join(os.path.expanduser("~"), "Downloads")
        output_file = os.path.join(downloads_folder, f"PrePost_Comparison_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

        try:
            generate_output(pre_file, post_file, output_file)
            messagebox.showinfo("Success", f"Output generated and saved to:\n{output_file}")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    tk.Button(root, text="Generate Output", bg="#2196F3", fg="white",
              font=("Arial", 12, "bold"), command=generate).pack(pady=18)

    root.mainloop()


if __name__ == "__main__":
    run_gui()
